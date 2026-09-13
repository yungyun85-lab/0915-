import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6] # completely blank layout

# Color Palette
BG_COLOR = RGBColor(246, 244, 238)         # Warm light cream/gray
CARD_BG = RGBColor(255, 255, 255)          # Pure white
BORDER_COLOR = RGBColor(229, 225, 216)     # Subtle border
TEXT_MAIN = RGBColor(45, 55, 72)           # Dark charcoal
TEXT_MUTED = RGBColor(107, 114, 128)       # Slate muted
TEXT_LIGHT = RGBColor(148, 163, 184)

SAGE_PRIMARY = RGBColor(91, 130, 102)      # #5B8266
SAGE_LIGHT = RGBColor(235, 242, 237)
TERRACOTTA = RGBColor(217, 125, 84)        # #D97D54
TERRA_LIGHT = RGBColor(250, 236, 231)
AMBER = RGBColor(232, 163, 61)             # #E8A33D
AMBER_LIGHT = RGBColor(253, 245, 232)
BLUE = RGBColor(2, 132, 199)               # #0284c7
BLUE_LIGHT = RGBColor(240, 249, 255)

FONT_HEADING = "Microsoft JhengHei"
FONT_BODY = "Microsoft JhengHei"

def add_header(slide, tag_text, title_text, subtitle_text=""):
    # Header tag pill
    if tag_text:
        tag_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.35)
        )
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = SAGE_LIGHT
        tag_box.line.color.rgb = SAGE_PRIMARY
        tag_box.line.width = Pt(1)
        tf = tag_box.text_frame
        tf.margin_top = Inches(0.04)
        p = tf.paragraphs[0]
        p.text = "● " + tag_text
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = SAGE_PRIMARY
        p.alignment = PP_ALIGN.CENTER
    
    # Title & Subtitle box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.95))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MAIN
    
    if subtitle_text:
        p3 = tf2.add_paragraph()
        p3.text = subtitle_text
        p3.font.name = FONT_BODY
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_MUTED

def set_slide_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

# ==========================================
# SLIDE 1: COVER
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

# Center main card
c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5))
c1.fill.solid()
c1.fill.fore_color.rgb = CARD_BG
c1.line.color.rgb = BORDER_COLOR
c1.line.width = Pt(1.5)

# Badge
pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(1.6), Inches(2.8), Inches(0.42))
pill.fill.solid()
pill.fill.fore_color.rgb = SAGE_LIGHT
pill.line.color.rgb = SAGE_PRIMARY
p = pill.text_frame.paragraphs[0]
p.text = "★ 數位教材創新成果專題"
p.font.name = FONT_HEADING
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = SAGE_PRIMARY
p.alignment = PP_ALIGN.CENTER

# Main Title Box
tb = s1.shapes.add_textbox(Inches(1.8), Inches(2.2), Inches(9.7), Inches(2.8))
tf = tb.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = 0

p1 = tf.paragraphs[0]
p1.text = "FocusLingua 專案成果匯報"
p1.font.name = FONT_HEADING
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.rgb = TEXT_MAIN

p2 = tf.add_paragraph()
p2.text = "希伯崙：AI英語互動學習設計 (FocusLingua 微型多感官學習對話系統)"
p2.font.name = FONT_HEADING
p2.font.size = Pt(18)
p2.font.bold = True
p2.font.color.rgb = SAGE_PRIMARY
p2.space_before = Pt(10)

p3 = tf.add_paragraph()
p3.text = "專為注意力缺失 (ADHD) 與學習障礙打造之 B2B2C 英語數位教材系統"
p3.font.name = FONT_BODY
p3.font.size = Pt(13)
p3.font.color.rgb = TEXT_MUTED
p3.space_before = Pt(8)

# Presenter Info Box
pbox = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(4.7), Inches(9.7), Inches(1.2))
pbox.fill.solid()
pbox.fill.fore_color.rgb = RGBColor(250, 248, 244)
pbox.line.color.rgb = BORDER_COLOR
ptf = pbox.text_frame
ptf.word_wrap = True
ptf.margin_left = Inches(0.3)
ptf.margin_top = Inches(0.18)

pp1 = ptf.paragraphs[0]
pp1.text = "發表者：陳詠芸 Anna Chen"
pp1.font.name = FONT_HEADING
pp1.font.size = Pt(14)
pp1.font.bold = True
pp1.font.color.rgb = TEXT_MAIN

pp2 = ptf.add_paragraph()
pp2.text = "中原大學 應用外語系主修 / 財務金融雙主修 ｜ IELTS 6.5 (L 7.5 / R 6.5)"
pp2.font.name = FONT_BODY
pp2.font.size = Pt(11)
pp2.font.color.rgb = TEXT_MUTED
pp2.space_before = Pt(4)


# ==========================================
# SLIDE 2: AGENDA
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "AGENDA", "簡報目錄 (Contents)", "五大核心進程，掌握現場洞察與產品落地")

agenda_items = [
    ("01", "自我介紹 (About Presenter)", "中原特資中心英語教學實務 × LiveABC 數位教材研發 × 跨領域背景", SAGE_PRIMARY),
    ("02", "專案動機 (Motivation)", "國小 1~6 年級課輔現場觀察 × 出題困境 × 數位教材公司現狀", TERRACOTTA),
    ("03", "FocusLingua 演進歷程 (Product Evolution)", "初步構思 ➔ 模型一試錯 ➔ 模型二轉向 ➔ 最終雙端解耦", BLUE),
    ("04", "雙端實體展示 (Demo Final)", "教師端 Web 智能備課工作站 × 學生端 iPhone 90 秒專注沙盒", AMBER),
    ("05", "自我收穫與合作展望 (Reflections & Partnership)", "敏銳觀察直擊真問題、與 AI 深度協作掌握減法、成就感與企業合作", SAGE_PRIMARY)
]

for idx, (num, title, desc, color) in enumerate(agenda_items):
    y_pos = Inches(1.9 + idx * 1.0)
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(0.85))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1)
    
    # Left number block
    num_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), y_pos + Inches(0.12), Inches(0.8), Inches(0.6))
    num_box.fill.solid()
    num_box.fill.fore_color.rgb = SAGE_LIGHT
    num_box.line.fill.background()
    np = num_box.text_frame.paragraphs[0]
    np.text = num
    np.font.name = FONT_HEADING
    np.font.size = Pt(16)
    np.font.bold = True
    np.font.color.rgb = color
    np.alignment = PP_ALIGN.CENTER
    
    # Text box
    tb = s2.shapes.add_textbox(Inches(1.9), y_pos + Inches(0.08), Inches(10.4), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(2)


# ==========================================
# SLIDE 3: SELF-INTRO
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "發表者背景", "自我介紹 (About Presenter)", "特資中心課輔實務 × 數位教材研發 × 人因工程與 AI 協作")

# Left Column: Profile Card
left_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(3.6), Inches(5.1))
left_card.fill.solid()
left_card.fill.fore_color.rgb = CARD_BG
left_card.line.color.rgb = BORDER_COLOR

# Photo
photo_path = "/Users/annamei/notes/assets/截圖 2026-09-09 晚上10.01.52.png"
if os.path.exists(photo_path):
    s3.shapes.add_picture(photo_path, Inches(1.1), Inches(2.1), Inches(3.0), Inches(2.3))

# Left Profile Info
ltb = s3.shapes.add_textbox(Inches(1.0), Inches(4.55), Inches(3.2), Inches(2.3))
ltf = ltb.text_frame
ltf.word_wrap = True
ltf.margin_left = ltf.margin_top = 0

lp1 = ltf.paragraphs[0]
lp1.text = "陳詠芸 Anna Chen"
lp1.font.name = FONT_HEADING
lp1.font.size = Pt(16)
lp1.font.bold = True
lp1.font.color.rgb = TEXT_MAIN
lp1.alignment = PP_ALIGN.CENTER

lp2 = ltf.add_paragraph()
lp2.text = "中原大學\n應用外語系主修 / 財務金融雙主修"
lp2.font.name = FONT_BODY
lp2.font.size = Pt(11)
lp2.font.color.rgb = TEXT_MUTED
lp2.alignment = PP_ALIGN.CENTER
lp2.space_before = Pt(4)

lp3 = ltf.add_paragraph()
lp3.text = "★ IELTS 6.5 (L 7.5 / R 6.5)"
lp3.font.name = FONT_HEADING
lp3.font.size = Pt(11)
lp3.font.bold = True
lp3.font.color.rgb = SAGE_PRIMARY
lp3.alignment = PP_ALIGN.CENTER
lp3.space_before = Pt(8)

# Right Column: 4 Experience Cards
exp_items = [
    ("中原特資中心 英語課輔老師", "特殊教育學習支持", 
     "長期於中原特資中心輔導特殊教育學生，第一線掌握注意力不集中 (ADHD) 與低專注耐受度之課堂認知極限；本專案 FocusLingua 進一步將此洞察轉化，以國小 1~6 年級為主要核心對象進行數位教材切片與鷹架系統研發。", TERRACOTTA),
    ("LiveABC 希伯崙 研發三處實習生 & 人事部工讀", "2026.01~02", 
     "參與英語數位教材企劃、教案研析與生成式 AI 切片；人事工讀累積跨部門協同能力。", SAGE_PRIMARY),
    ("加拿大 Athabasca Univ. VIP Research 實習", "2025.07~08", 
     "使用 MEGA World 開發英語教學遊戲，融合 TPR (全身肢體反應) 與 CLT (溝通教學法)。", BLUE),
    ("2025 全國前瞻科技英語教學創新競賽", "優等第二名", 
     "以創新互動教學技術獲全國評審肯定，兼具教育科技整合與實作落地能力。", AMBER)
]

for idx, (title, tag, desc, color) in enumerate(exp_items):
    y_pos = Inches(1.9 + idx * 1.25)
    ec = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.7), y_pos, Inches(7.833), Inches(1.15))
    ec.fill.solid()
    ec.fill.fore_color.rgb = CARD_BG
    ec.line.color.rgb = BORDER_COLOR
    ec.line.width = Pt(1)
    
    # Left accent strip
    strip = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.7), y_pos, Inches(0.12), Inches(1.15))
    strip.fill.solid()
    strip.fill.fore_color.rgb = color
    strip.line.fill.background()
    
    etb = s3.shapes.add_textbox(Inches(4.95), y_pos + Inches(0.08), Inches(7.4), Inches(1.0))
    etf = etb.text_frame
    etf.word_wrap = True
    etf.margin_left = etf.margin_top = 0
    
    ep1 = etf.paragraphs[0]
    ep1.text = f"{title}   [{tag}]"
    ep1.font.name = FONT_HEADING
    ep1.font.size = Pt(12)
    ep1.font.bold = True
    ep1.font.color.rgb = TEXT_MAIN
    
    ep2 = etf.add_paragraph()
    ep2.text = desc
    ep2.font.name = FONT_BODY
    ep2.font.size = Pt(9.5)
    ep2.font.color.rgb = TEXT_MUTED
    ep2.space_before = Pt(3)


# ==========================================
# SLIDE 4: MOTIVATION
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "現場洞察", "專案動機 (Motivation)", "國小 1~6 年級課輔現場痛點 × 數位教材出版業現狀（概念提案，非臨床實驗）")

cards_data = [
    ("學生端痛點 (國小 1~6 年級)", [
        "學習時間多花 2 倍：注意力不集中 (ADHD) 啟動阻力大、極易疲乏分心。",
        "極度抗拒課後複習：回家缺乏結構化環境，強烈傾向在課堂內直接學會。"
    ], TERRACOTTA),
    ("課輔老師困境", [
        "手邊只有一本課本與課綱：缺乏針對注意力不集中學生的分層教材。",
        "不知原任課老師出題風格：資訊不對稱，時間緊迫下備課困難，更難手動出題。"
    ], SAGE_PRIMARY),
    ("公司現狀 (LiveABC)", [
        "傳統題庫一體適用：15~20 分鐘連貫長關卡，缺乏適合國小生的 90 秒微切片。",
        "缺乏特教評估數據：僅有單一正答率，無法產出符合學校需求之 IEP 行為進程指標。"
    ], BLUE)
]

for idx, (title, points, color) in enumerate(cards_data):
    x_pos = Inches(0.8 + idx * 4.0)
    box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.9), Inches(3.733), Inches(3.6))
    box.fill.solid()
    box.fill.fore_color.rgb = CARD_BG
    box.line.color.rgb = BORDER_COLOR
    
    # Top accent line
    top_line = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_pos, Inches(1.9), Inches(3.733), Inches(0.1))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = color
    top_line.line.fill.background()
    
    tb = s4.shapes.add_textbox(x_pos + Inches(0.2), Inches(2.15), Inches(3.333), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    
    for pt in points:
        pp = tf.add_paragraph()
        pp.text = "• " + pt
        pp.font.name = FONT_BODY
        pp.font.size = Pt(10.5)
        pp.font.color.rgb = TEXT_MUTED
        pp.space_before = Pt(10)

# Bottom Key Proposition Callout
bot = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.75), Inches(11.733), Inches(1.2))
bot.fill.solid()
bot.fill.fore_color.rgb = SAGE_LIGHT
bot.line.color.rgb = SAGE_PRIMARY
btf = bot.text_frame
btf.margin_left = Inches(0.3)
btf.margin_top = Inches(0.15)
bp1 = btf.paragraphs[0]
bp1.text = "💡 破局核心命題："
bp1.font.name = FONT_HEADING
bp1.font.size = Pt(13)
bp1.font.bold = True
bp1.font.color.rgb = SAGE_PRIMARY
bp2 = btf.add_paragraph()
bp2.text = "「能不能用 AI 在 30 秒內替老師自動切片國小課本與出題，讓國小 1~6 年級孩子在課堂 90 秒內無痛完成微任務，同時自動沉澱特教 IEP 評估數據？」"
bp2.font.name = FONT_BODY
bp2.font.size = Pt(11)
bp2.font.color.rgb = TEXT_MAIN
bp2.space_before = Pt(4)


# ==========================================
# SLIDE 5: EVOLUTION (4 STAGES)
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "產品演進脈絡", "FocusLingua：從發現問題到系統化落地的四階段演進", "不盲信技術，經歷兩代失敗原型的深刻反思，淬鍊出真正的 B2B2C 雙端解耦架構")

stages = [
    ("階段一", "初步構思 (Initial Idea)", [
        "出發點：專為國小 1~6 年級 ADHD 注意力不集中族群構思英語學習 App。",
        "核心目標：解決國小孩子背單字耗時比別人長（2倍以上）且課堂極易分心的問題。",
        "局限：初期僅有概念發想，缺乏第一線教學現場的具體載體。"
    ], RGBColor(148, 163, 184)),
    ("模型一", "AI 原型摸索 (Stitch/Studio)", [
        "實作：使用 Agent, Claude, Stitch, Studio AI 快速拼出第一版雛形。",
        "失敗洞察：\n1. 介面僵硬、互動不自然。\n2. 未切分低中高年級，繁複選單讓學童分心！",
        "教訓：注意力不集中孩子需要直覺引導，選項過多是認知干擾。"
    ], TERRACOTTA),
    ("模型二", "Chatbot 轉向與老師端萌芽", [
        "嘗試：引入首頁 Chatbot 對話式引導，試圖以問答輔助學童。",
        "失敗洞察：\n1. 互動自由度過高，學童偏離學習。\n2. Token 與伺服器成本高昂且延遲難控。\n3. 缺乏課輔老師把關。",
        "教訓：AI 不應直接放任無邊界互動，應聚焦於繁重教材的減法切片。"
    ], AMBER),
    ("最終落地", "雙端解耦與 B2B2C 架構", [
        "教師端 Web：AI 30 秒切片 ＋ 人工二次修改把關 ＋ 一鍵派發。",
        "學生端 App：90 秒沙漏微任務 ＋ 多巴胺即時反饋，零分心設計。",
        "成功解法：以課輔老師為軸心，AI 作為減法工具，完美契合現場節奏。"
    ], SAGE_PRIMARY)
]

for idx, (badge, title, pts, color) in enumerate(stages):
    x_pos = Inches(0.8 + idx * 2.98)
    card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.9), Inches(2.8), Inches(5.1))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    
    # Top strip
    strip = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_pos, Inches(1.9), Inches(2.8), Inches(0.1))
    strip.fill.solid()
    strip.fill.fore_color.rgb = color
    strip.line.fill.background()
    
    tb = s5.shapes.add_textbox(x_pos + Inches(0.15), Inches(2.1), Inches(2.5), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    p.text = badge
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = color
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MAIN
    p2.space_before = Pt(4)
    
    for pt in pts:
        p3 = tf.add_paragraph()
        p3.text = pt
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9.2)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(6)


# ==========================================
# SLIDE 6: DEMO FINAL
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "實體演示", "教師端 Web 工作站 ＋ 學生端 App 雙端即時交互展示", "左側：國小教材上傳 ➔ AI 30 秒切片 ➔ 二次修改 ➔ 派發 ｜ 右側：國小學生手機 90 秒微任務體驗")

# Left Column: Teacher Web Box
tw = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(6.8), Inches(5.1))
tw.fill.solid()
tw.fill.fore_color.rgb = CARD_BG
tw.line.color.rgb = BORDER_COLOR
ttb = s6.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(6.2), Inches(4.7))
ttf = ttb.text_frame
ttf.word_wrap = True
ttf.margin_left = ttf.margin_top = 0

tp1 = ttf.paragraphs[0]
tp1.text = "💻 教師端備課工作站 (Desktop Web)"
tp1.font.name = FONT_HEADING
tp1.font.size = Pt(15)
tp1.font.bold = True
tp1.font.color.rgb = SAGE_PRIMARY

t_points = [
    ("1. 國小英語教材上傳與 AI 切片", "支援課綱 PDF/PPT 與聽力音檔，AI 於 30 秒內自動拆解為低、中、高年級之 90 秒微任務題型。"),
    ("2. 題目二次修改與人工審核台", "老師能直接微調句子長度、增刪鷹架輔助提示詞，把關教學品質後一鍵派發至班級。"),
    ("3. 班級注意力與行為數據追蹤", "即時記錄全班學童答題延遲 (Latency)、猶豫拐點及衝動性亂點次數，自動輸出標準 IEP 報表。")
]
for t_tit, t_desc in t_points:
    p = ttf.add_paragraph()
    p.text = "• " + t_tit
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)
    
    p2 = ttf.add_paragraph()
    p2.text = t_desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(2)

# Right Column: Student App Box
sw = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.9), Inches(4.733), Inches(5.1))
sw.fill.solid()
sw.fill.fore_color.rgb = CARD_BG
sw.line.color.rgb = BORDER_COLOR
stb = s6.shapes.add_textbox(Inches(8.1), Inches(2.1), Inches(4.133), Inches(4.7))
stf = stb.text_frame
stf.word_wrap = True
stf.margin_left = stf.margin_top = 0

sp1 = stf.paragraphs[0]
sp1.text = "📱 學生端手機 App 專注沙盒 (Mobile Sandbox)"
sp1.font.name = FONT_HEADING
sp1.font.size = Pt(15)
sp1.font.bold = True
sp1.font.color.rgb = TERRACOTTA

s_points = [
    ("1. 90 秒沙漏倒數計時", "將大單元粉碎為 90 秒微型挑戰，消除 ADHD 學童對冗長課程的時間焦慮。"),
    ("2. 多感官伴讀卡與音節切片", "Lexend 友善字型切換、高對比顏色切片與原生單字真人發音輔助。"),
    ("3. 動態積木拼句與多巴胺即時反饋", "捨棄傳統鍵盤手打，以點擊積木拼裝降低挫折，答對即刻獲得觸覺與微粒子視覺獎勵。"),
    ("4. 課堂防沉迷護眼鎖定機制", "完成微任務後啟動防沉迷鎖定，引導閉眼休息 10 分鐘，保護專注神經。")
]
for s_tit, s_desc in s_points:
    p = stf.add_paragraph()
    p.text = "• " + s_tit
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(6)
    
    p2 = stf.add_paragraph()
    p2.text = s_desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(1)


# ==========================================
# SLIDE 7: REFLECTIONS & COLLABORATION
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "收穫與展望", "自我收穫與合作展望 (Reflections & Partnership)", "從第一線痛點出發，與 AI 深度協作迭代解法，成就感與企業合作願景")

r_cards = [
    ("1. 敏銳觀察，直擊課堂真問題", [
        "不再流於空想技術：如果沒有親自站在課輔講台，永遠不知道注意力不集中的孩子有多排斥課後作業，更不會知道課輔老師手握一本課本卻不知如何出題的巨大無助。",
        "教育同理心驅動：專案最大的意義，在於讓「慢飛」不是孩子的原罪，而是用合適的科技工具撫平學習門檻。"
    ], SAGE_PRIMARY),
    ("2. 與 AI 深度協作，掌握減法設計", [
        "從挫折中淬鍊：經歷模型一（自選過多反而分心）與模型二（Chatbot 昂貴低效）的試錯，我深刻理解到 AI 最佳的定位不是無差別炫技，而是「繁重切片的減法工具」。",
        "人機平衡：讓 AI 負責 30 秒自動消化課綱出題，讓人（老師）把關二次審核，實現可靠、低成本的落地閉環。"
    ], TERRACOTTA),
    ("3. 企業合作願景 (LiveABC 希伯崙落地延伸)", [
        "B2B2C 商業閉環：出版業（LiveABC）提供教材授權與切片模組 ➔ 學校/課輔機構採購教師工作站 ➔ 學生課堂高效落實。",
        "打造普惠特教標竿：為傳統出版社開拓特殊教育新藍海，實現商業價值與社會公益雙贏。"
    ], BLUE)
]

for idx, (title, pts, color) in enumerate(r_cards):
    y_pos = Inches(1.9 + idx * 1.65)
    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(1.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1)
    
    # Left strip
    strip = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(0.15), Inches(1.5))
    strip.fill.solid()
    strip.fill.fore_color.rgb = color
    strip.line.fill.background()
    
    tb = s7.shapes.add_textbox(Inches(1.2), y_pos + Inches(0.12), Inches(11.1), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    
    for pt in pts:
        p2 = tf.add_paragraph()
        p2.text = "• " + pt
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.2)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

# Output Paths
desktop_path = "/Users/annamei/Desktop/FocusLingua_成果專題簡報.pptx"
notes_path = "/Users/annamei/notes/FocusLingua_成果專題簡報.pptx"

prs.save(desktop_path)
prs.save(notes_path)
print(f"PPTX successfully generated and saved to:\n1. {desktop_path}\n2. {notes_path}")
